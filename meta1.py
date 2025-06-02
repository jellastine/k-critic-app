import streamlit as st
import random
import os
import joblib
import pandas as pd
import matplotlib.pyplot as plt
from PyPDF2 import PdfReader
from docx import Document
from io import BytesIO
from fpdf import FPDF
from wordcloud import WordCloud

# --- 문서 요약 (임시 요약 텍스트)
def summarize_text(text):
    return text[:500] + '...'

# --- 파일 파싱
@st.cache_data
def extract_text_from_file(uploaded_file):
    if uploaded_file.name.endswith(".pdf"):
        reader = PdfReader(uploaded_file)
        return "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
    elif uploaded_file.name.endswith(".docx"):
        doc = Document(uploaded_file)
        return "\n".join([para.text for para in doc.paragraphs])
    elif uploaded_file.name.endswith(".txt"):
        return uploaded_file.read().decode("utf-8")
    else:
        return "지원하지 않는 파일 형식입니다."

# --- 모델 불러오기 (학습된 모델 필요)
@st.cache_resource
def load_model():
    return joblib.load("metacritic_predictor.pkl")

# --- PDF 저장 함수
def save_report_to_pdf(game_name, predicted_score, summary, unique_features, main_elements, competitors):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt=f"K-크리틱 분석 리포트 - {game_name}", ln=True, align='C')
    pdf.ln(10)
    pdf.multi_cell(0, 10, txt=f"예상 메타크리틱 점수: {predicted_score:.1f}\n\n문서 요약:\n{summary}\n\n강점 요약:\n- 차별화 포인트: {unique_features}\n- 핵심 시스템: {main_elements}\n\n경쟁작:\n{competitors}")
    buffer = BytesIO()
    pdf.output(buffer)
    return buffer

# --- Streamlit UI
st.title("🎮 K-크리틱: 게임 기획 AI 평가 시스템 (MVP)")

st.header("1. 게임 기획 문서 업로드")
uploaded_file = st.file_uploader("기획 문서를 업로드하세요 (PDF, Word, 텍스트)", type=["pdf", "docx", "txt"])

st.header("2. 게임 기본 정보 입력")
with st.form("input_form"):
    game_name = st.text_input("게임 이름")
    genre = st.text_input("장르")
    platform = st.selectbox("출시 플랫폼", ["PC", "PS5", "Xbox", "Switch", "Mobile"])
    play_mode = st.radio("플레이 방식", ["싱글", "멀티", "싱글+멀티"])
    play_time = st.slider("예상 플레이 타임 (시간)", 1, 100, 10)
    user_score = st.slider("예상 유저 점수 (1.0 ~ 10.0)", 1.0, 10.0, 7.5)
    release_year = st.number_input("출시 연도", min_value=2000, max_value=2030, value=2025)
    target_audience = st.text_input("타겟 유저층")
    competitors = st.text_input("경쟁작 (쉼표로 구분)")
    unique_features = st.text_area("차별화 포인트")
    main_elements = st.text_area("핵심 게임 요소")
    submitted = st.form_submit_button("K-크리틱 분석 시작")

if submitted:
    if not uploaded_file:
        st.warning("⚠️ 기획 문서를 먼저 업로드해주세요.")
    else:
        with st.spinner("문서 분석 및 점수 예측 중..."):
            text = extract_text_from_file(uploaded_file)
            summary = summarize_text(text)

            # 모델 로드 및 예측
            model = load_model()
            input_df = pd.DataFrame([{
                "genre": genre,
                "platform": platform,
                "user_score": user_score,
                "release_year": release_year
            }])
            predicted_score = model.predict(input_df)[0]

        st.success("분석 완료!")

        st.header("📊 K-크리틱 평가 보고서")
        st.subheader(f"예상 메타크리틱 점수: **{predicted_score:.1f}점**")

        st.subheader("문서 요약")
        st.write(summary)

        st.subheader("강점 요약")
        st.markdown(f"- 차별화 포인트: {unique_features}")
        st.markdown(f"- 핵심 시스템: {main_elements}")

        st.subheader("경쟁작")
        st.write(competitors)

        # --- 워드클라우드 시각화 ---
        st.subheader("📌 핵심 키워드 시각화")
        all_text = f"{summary} {unique_features} {main_elements}"
        wordcloud = WordCloud(width=800, height=400, background_color='white').generate(all_text)
        fig, ax = plt.subplots(figsize=(10, 5))
        ax.imshow(wordcloud, interpolation='bilinear')
        ax.axis('off')
        st.pyplot(fig)

        # --- PDF 다운로드 버튼 ---
        st.subheader("📄 PDF 보고서 저장")
        pdf_buffer = save_report_to_pdf(game_name, predicted_score, summary, unique_features, main_elements, competitors)
        st.download_button(label="📥 PDF 다운로드", data=pdf_buffer.getvalue(), file_name=f"{game_name}_report.pdf", mime="application/pdf")

        st.info("✅ PDF 저장, 시각화, 요약 기능 통합 완료. 향후 GPT 분석 추가 예정")
